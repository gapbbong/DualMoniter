import sys
import asyncio
import threading
if sys.platform == "win32":
    asyncio.set_event_loop_policy(asyncio.WindowsProactorEventLoopPolicy())
    import asyncio.proactor_events
    import asyncio.windows_events
    import asyncio.trsock
    from asyncio import exceptions, tasks
    import _overlapped
    import struct
    import socket

    # Patch 1: IocpProactor.accept - handle client disconnection during handshake
    _orig_accept = asyncio.windows_events.IocpProactor.accept
    def _safe_accept(self, listener):
        self._register_with_iocp(listener)
        conn = self._get_accept_socket(listener.family)
        ov = _overlapped.Overlapped(asyncio.windows_events.NULL)
        ov.AcceptEx(listener.fileno(), conn.fileno())

        def finish_accept(trans, key, ov):
            ov.getresult()
            buf = struct.pack('@P', listener.fileno())
            conn.setsockopt(socket.SOL_SOCKET, _overlapped.SO_UPDATE_ACCEPT_CONTEXT, buf)
            conn.settimeout(listener.gettimeout())
            return conn, conn.getpeername()

        async def accept_coro(future, conn):
            try:
                await future
            except (exceptions.CancelledError, OSError):
                try:
                    conn.close()
                except Exception:
                    pass
                if not isinstance(sys.exc_info()[1], OSError):
                    raise

        future = self._register(ov, listener, finish_accept)
        coro = accept_coro(future, conn)
        tasks.ensure_future(coro, loop=self._loop)
        return future

    asyncio.windows_events.IocpProactor.accept = _safe_accept

    # Patch 2: BaseProactorEventLoop._start_serving - prevent server socket closure on client transient resets
    _orig_start_serving = asyncio.proactor_events.BaseProactorEventLoop._start_serving

    def _safe_start_serving(self, protocol_factory, sock,
                            sslcontext=None, server=None, backlog=100,
                            ssl_handshake_timeout=None,
                            ssl_shutdown_timeout=None):
        def loop(f=None):
            try:
                if f is not None:
                    try:
                        conn, addr = f.result()
                    except OSError as exc:
                        if getattr(exc, 'winerror', None) in (64, 121, 995, 1225, 10054) or getattr(exc, 'errno', None) in (64, 121, 995, 1225, 10054):
                            if not self.is_closed() and sock.fileno() != -1:
                                f = self._proactor.accept(sock)
                                self._accept_futures[sock.fileno()] = f
                                f.add_done_callback(loop)
                            return
                        raise
                    if self._debug:
                        asyncio.proactor_events.logger.debug("%r got a new connection from %r: %r",
                                     server, addr, conn)
                    protocol = protocol_factory()
                    if sslcontext is not None:
                        self._make_ssl_transport(
                            conn, protocol, sslcontext, server_side=True,
                            extra={'peername': addr}, server=server,
                            ssl_handshake_timeout=ssl_handshake_timeout,
                            ssl_shutdown_timeout=ssl_shutdown_timeout)
                    else:
                        self._make_socket_transport(
                            conn, protocol,
                            extra={'peername': addr}, server=server)
                if self.is_closed():
                    return
                f = self._proactor.accept(sock)
            except OSError as exc:
                if sock.fileno() != -1:
                    self.call_exception_handler({
                        'message': 'Accept failed on a socket',
                        'exception': exc,
                        'socket': asyncio.trsock.TransportSocket(sock),
                    })
                    sock.close()
                elif self._debug:
                    asyncio.proactor_events.logger.debug("Accept failed on socket %r",
                                 sock, exc_info=True)
            except asyncio.exceptions.CancelledError:
                sock.close()
            else:
                self._accept_futures[sock.fileno()] = f
                f.add_done_callback(loop)

        self.call_soon(loop)

    asyncio.proactor_events.BaseProactorEventLoop._start_serving = _safe_start_serving

from fastapi import FastAPI, WebSocket, WebSocketDisconnect, UploadFile, File, Form, Request
from fastapi.staticfiles import StaticFiles
from fastapi.responses import HTMLResponse, JSONResponse, FileResponse
from typing import Dict, List
import uvicorn
import os
import json
import time
import socket
import base64
import requests
import re

IMAGE_DIR = "학생_화면_저장폴더"
if not os.path.exists(IMAGE_DIR):
    os.makedirs(IMAGE_DIR)

# 대시보드 캡처 저장 폴더 (C:\Users\user\Pictures\CIM 대시보드)
CAPTURE_DIR = os.path.join(os.environ.get("USERPROFILE", r"C:\Users\user"), "Pictures", "CIM 대시보드")
if not os.path.exists(CAPTURE_DIR):
    try:
        os.makedirs(CAPTURE_DIR, exist_ok=True)
    except Exception as e:
        print(f"[Error] CAPTURE_DIR creation: {e}")

app = FastAPI()
app.mount("/static", StaticFiles(directory="static"), name="static")
from fastapi.responses import HTMLResponse, JSONResponse, Response

@app.get("/screenshots/{filename}")
async def get_screenshot(filename: str):
    file_path = os.path.join(IMAGE_DIR, filename)
    if os.path.exists(file_path):
        try:
            with open(file_path, "rb") as f:
                content = f.read()
            return Response(content=content, media_type="image/jpeg")
        except:
            pass
    return Response(status_code=404)

# 연결된 클라이언트 웹소켓 관리
active_websockets: Dict[str, WebSocket] = {}

LOG_HISTORY_FILE = "all_agent_logs.json"
if not os.path.exists(LOG_HISTORY_FILE):
    with open(LOG_HISTORY_FILE, "w", encoding="utf-8") as f:
        json.dump([], f)

BLOCKLIST_FILE = "blocklist.json"
if not os.path.exists(BLOCKLIST_FILE):
    with open(BLOCKLIST_FILE, "w", encoding="utf-8") as f:
        json.dump([], f)

# client_id(PC 단위) -> 최근 보고된 실행 프로세스 이름 목록
client_processes: Dict[str, list] = {}

@app.get("/", response_class=HTMLResponse)
async def get_dashboard():
    with open("index.html", "r", encoding="utf-8") as f:
        return f.read()

@app.get("/submit", response_class=HTMLResponse)
async def get_submit_page():
    if os.path.exists("submit.html"):
        with open("submit.html", "r", encoding="utf-8") as f:
            return f.read()
    return "<h3>과제 제출 페이지를 찾을 수 없습니다.</h3>"

@app.get("/download_agent")
async def download_agent():
    from fastapi.responses import FileResponse
    agent_path = "FINAL_AGENT_PRO.py"
    if os.path.exists(agent_path):
        return FileResponse(agent_path, media_type="application/octet-stream", filename="FINAL_AGENT_PRO.py")
    from fastapi.responses import JSONResponse
    return JSONResponse(status_code=404, content={"message": "Agent file not found"})

@app.get("/check_update")
async def check_update(hash: str = ""):
    import hashlib
    agent_path = "FINAL_AGENT_PRO.py"
    if not os.path.exists(agent_path):
        return {"update_needed": False}
    try:
        with open(agent_path, "rb") as f:
            raw_bytes = f.read()
            server_hash = hashlib.md5(raw_bytes).hexdigest()
        win_hash = hashlib.md5(raw_bytes.replace(b"\r\n", b"\n").replace(b"\n", b"\r\n")).hexdigest()
        unix_hash = hashlib.md5(raw_bytes.replace(b"\r\n", b"\n")).hexdigest()
        valid_hashes = {server_hash, win_hash, unix_hash}
        needed = (hash not in valid_hashes)
        return {"update_needed": needed}
    except:
        return {"update_needed": False}

@app.websocket("/ws/{client_id}")
async def websocket_endpoint(websocket: WebSocket, client_id: str):
    await websocket.accept()
    active_websockets[client_id] = websocket
    print(f"[WS] Client connected: {client_id}")
    try:
        while True:
            # 클라이언트로부터 메시지 수신 (심박수 체크용)
            data = await websocket.receive_text()
            # 필요 시 메시지 처리
    except WebSocketDisconnect:
        if client_id in active_websockets:
            del active_websockets[client_id]
        print(f"[WS] Client disconnected: {client_id}")

@app.post("/log/{client_id}")
async def receive_log(client_id: str, request: Request):
    data = await request.json()
    message = data.get("message")
    log_entry = f"[{time.strftime('%H:%M:%S')}] [{client_id}] {message}"
    try:
        with open(LOG_HISTORY_FILE, "r+", encoding="utf-8") as f:
            logs = json.load(f)
            logs.insert(0, log_entry)
            f.seek(0)
            json.dump(logs[:50], f, ensure_ascii=False)
            f.truncate()
    except: pass
    return {"status": "success"}

def get_room_prefix():
    # 1. room_config.txt 파일이 있으면 우선 사용 (예: prefix=204 또는 204)
    if os.path.exists("room_config.txt"):
        try:
            with open("room_config.txt", "r", encoding="utf-8") as f:
                for line in f:
                    line = line.strip()
                    if line and not line.startswith("#") and "prefix=" in line:
                        return line.split("prefix=")[1].strip()
                    elif line and not line.startswith("#") and not "=" in line:
                        return line
        except: pass

    # 2. 교사 PC 컴퓨터 이름에서 자동 추출 (예: 204-TEACHER -> 204)
    try:
        hn = socket.gethostname()
        parts = hn.split("-")
        if len(parts) >= 2 and parts[0].isdigit():
            return parts[0]
    except: pass

    return "203" # 기본값

def load_student_names():
    names_map = {}
    try:
        import openpyxl
        # 1. room_config.txt에 지정된 엑셀 파일명 확인 (예: excel=2-1반.xlsx)
        target_xlsx = None
        if os.path.exists("room_config.txt"):
            try:
                with open("room_config.txt", "r", encoding="utf-8") as f:
                    for line in f:
                        if "excel=" in line:
                            target_xlsx = line.split("excel=")[1].strip()
            except: pass

        if not target_xlsx or not os.path.exists(target_xlsx):
            # 2. 폴더 내의 첫 번째 .xlsx 파일 자동 탐색 (또는 3-5반.xlsx)
            xlsx_files = [f for f in os.listdir(".") if f.endswith(".xlsx") and not f.startswith("~")]
            target_xlsx = xlsx_files[0] if xlsx_files else "3-5반.xlsx"

        if target_xlsx and os.path.exists(target_xlsx):
            wb = openpyxl.load_workbook(target_xlsx, data_only=True)
            sheet = wb.active
            for r in range(2, sheet.max_row + 1):
                pc_val = sheet.cell(row=r, column=1).value
                name_val = sheet.cell(row=r, column=2).value
                if pc_val:
                    pc_str = str(pc_val).strip()
                    name_str = str(name_val).strip() if name_val else ""
                    names_map[pc_str] = name_str
    except Exception as e:
        print(f"[Error] Failed to load student names: {e}")
    return names_map


def get_local_server_ip():
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.connect(("8.8.8.8", 80))
        ip = s.getsockname()[0]
        s.close()
        return ip
    except Exception:
        pass
    try:
        ips = socket.gethostbyname_ex(socket.gethostname())[2]
        for ip in ips:
            if not ip.startswith("127.") and not ip.startswith("169.254."):
                return ip
    except Exception:
        pass
    return "127.0.0.1"

DISCOVERY_PORT = 8889

def start_discovery_server():
    def discovery_worker():
        sock = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        sock.setsockopt(socket.SOL_SOCKET, socket.SO_REUSEADDR, 1)
        sock.setsockopt(socket.SOL_SOCKET, socket.SO_BROADCAST, 1)
        try:
            sock.bind(("", DISCOVERY_PORT))
        except Exception as e:
            print(f"[Discovery] Port {DISCOVERY_PORT} bind warning: {e}")
            return

        sock.settimeout(1.0)
        last_beacon = 0

        while True:
            now = time.time()
            server_ip = get_local_server_ip()
            prefix = get_room_prefix()

            # 3초마다 LAN 브로드캐스트 비콘 전송
            if now - last_beacon > 3.0:
                last_beacon = now
                beacon_msg = f"CIM_BEACON|{server_ip}|8888|{prefix}".encode("utf-8")
                try:
                    sock.sendto(beacon_msg, ("255.255.255.255", DISCOVERY_PORT))
                except Exception:
                    pass

            # 학생 PC의 탐색 요청(DISCOVER) 수신 시 즉각 응답
            try:
                data, addr = sock.recvfrom(1024)
                text = data.decode("utf-8", errors="ignore").strip()
                if text.startswith("CIM_DISCOVER"):
                    req_parts = text.split("|")
                    req_room = req_parts[1].strip() if len(req_parts) >= 2 else ""
                    # 요청한 실습실 번호가 있고 현재 교사 실습실과 다르면 다른 교실이므로 무시
                    if req_room and prefix and req_room != prefix:
                        pass
                    else:
                        resp_msg = f"CIM_RESPONSE|{server_ip}|8888|{prefix}".encode("utf-8")
                        sock.sendto(resp_msg, addr)
            except socket.timeout:
                pass
            except Exception:
                pass

    t = threading.Thread(target=discovery_worker, daemon=True)
    t.start()
    print("[Discovery] LAN Zero-Config Auto-Discovery service running on UDP 8889")

@app.get("/get_config")
async def get_config():
    prefix = get_room_prefix()
    names = load_student_names()
    return JSONResponse(content={"prefix": prefix, "names": names})

active_windows: Dict[str, str] = {}

@app.get("/get_status")
async def get_status():
    prefix = get_room_prefix()
    status = {}
    now = time.time()
    for i in range(1, 25):
        base_id = f"{prefix}-{i:02d}"
        for suffix in ["1", "2"]:
            client_id = f"{base_id}-{suffix}"
            file_path = os.path.join(IMAGE_DIR, f"screen_{client_id}.jpg")
            is_online = False
            mtime_str = "-"
            if os.path.exists(file_path):
                mtime = os.path.getmtime(file_path)
                is_online = (now - mtime) < 20
                mtime_str = time.strftime("%H:%M:%S", time.localtime(mtime))
            
            status[client_id] = {
                "online": is_online,
                "time": mtime_str,
                "title": active_windows.get(client_id, "")
            }
    # 추가로 업로드된 클라이언트 화면 자동 포함
    try:
        for fname in os.listdir(IMAGE_DIR):
            if fname.startswith("screen_") and fname.endswith(".jpg"):
                cid = fname[7:-4]
                if cid not in status:
                    fpath = os.path.join(IMAGE_DIR, fname)
                    mtime = os.path.getmtime(fpath)
                    status[cid] = {
                        "online": (now - mtime) < 20,
                        "time": time.strftime("%H:%M:%S", time.localtime(mtime)),
                        "title": active_windows.get(cid, "")
                    }
    except Exception:
        pass
    return JSONResponse(content=status)

@app.get("/get_logs")
async def get_logs():
    try:
        with open(LOG_HISTORY_FILE, "r", encoding="utf-8") as f:
            return JSONResponse(content=json.load(f))
    except: return JSONResponse(content=[])

# ===== 학생 PC 사용 기록 (타임라인 & 1분 무반응 감지) =====
TIMELINE_FILE = "all_timelines.json"
client_timelines: Dict[str, dict] = {}

def load_timelines():
    global client_timelines
    if os.path.exists(TIMELINE_FILE):
        try:
            with open(TIMELINE_FILE, "r", encoding="utf-8") as f:
                client_timelines = json.load(f)
        except Exception as e:
            print(f"[Timeline] Load failed: {e}")

def save_timelines():
    try:
        with open(TIMELINE_FILE, "w", encoding="utf-8") as f:
            json.dump(client_timelines, f, ensure_ascii=False)
    except Exception as e:
        print(f"[Timeline] Save failed: {e}")

load_timelines()

def format_duration(seconds: int) -> str:
    if seconds < 60:
        return f"{seconds}초"
    elif seconds < 3600:
        m = seconds // 60
        s = seconds % 60
        return f"{m}분 {s}초" if s > 0 else f"{m}분"
    else:
        h = seconds // 3600
        m = (seconds % 3600) // 60
        return f"{h}시간 {m}분"

def update_client_timeline(base_id: str, title: str, idle_sec: int):
    now_ts = time.time()
    now_str = time.strftime("%H:%M:%S")
    is_idle = (idle_sec >= 60)
    
    clean_title = title.strip() if title else ""
    if is_idle:
        display_title = "아무 반응 없음"
    else:
        display_title = clean_title if clean_title else "바탕화면 / 파일 탐색기"

    if base_id not in client_timelines:
        client_timelines[base_id] = {
            "boot_time": now_str,
            "first_seen_ts": now_ts,
            "last_seen_ts": now_ts,
            "current_state": {
                "title": display_title,
                "is_idle": is_idle,
                "idle_sec": idle_sec,
                "start_time": now_str,
                "end_time": now_str,
                "start_ts": now_ts,
                "end_ts": now_ts,
                "duration_sec": 1,
                "duration_str": "1초"
            },
            "history": []
        }
        save_timelines()
        return

    data = client_timelines[base_id]
    data["last_seen_ts"] = now_ts
    curr = data.get("current_state")

    if not curr:
        data["current_state"] = {
            "title": display_title,
            "is_idle": is_idle,
            "idle_sec": idle_sec,
            "start_time": now_str,
            "end_time": now_str,
            "start_ts": now_ts,
            "end_ts": now_ts,
            "duration_sec": 1,
            "duration_str": "1초"
        }
        save_timelines()
        return

    # 상태 변경 여부 판단
    state_changed = False
    if curr.get("is_idle") != is_idle:
        state_changed = True
    elif not is_idle and curr.get("title") != display_title:
        state_changed = True

    if not state_changed:
        # 동일 상태 유지 -> 종료 시간 및 지속 시간 갱신
        curr["end_time"] = now_str
        curr["end_ts"] = now_ts
        curr["idle_sec"] = idle_sec
        dur = max(1, int(now_ts - curr.get("start_ts", now_ts)))
        curr["duration_sec"] = dur
        curr["duration_str"] = format_duration(dur)
    else:
        # 상태 전환됨 -> 기존 상태 확정 후 history에 보관, 새 current_state 시작
        curr["end_time"] = now_str
        curr["end_ts"] = now_ts
        dur = max(1, int(now_ts - curr.get("start_ts", now_ts)))
        curr["duration_sec"] = dur
        curr["duration_str"] = format_duration(dur)

        if "history" not in data:
            data["history"] = []
        data["history"].insert(0, dict(curr))
        if len(data["history"]) > 500:
            data["history"] = data["history"][:500]

        data["current_state"] = {
            "title": display_title,
            "is_idle": is_idle,
            "idle_sec": idle_sec,
            "start_time": now_str,
            "end_time": now_str,
            "start_ts": now_ts,
            "end_ts": now_ts,
            "duration_sec": 1,
            "duration_str": "1초"
        }
        save_timelines()

@app.get("/get_timeline/{base_id}")
async def get_timeline(base_id: str):
    data = client_timelines.get(base_id)
    if not data:
        return JSONResponse(content={
            "base_id": base_id,
            "boot_time": "-",
            "is_online": False,
            "total_records": 0,
            "timeline": []
        })

    timeline_list = []
    curr = data.get("current_state")
    now_ts = time.time()
    if curr:
        curr_copy = dict(curr)
        curr_copy["end_time"] = time.strftime("%H:%M:%S")
        dur = max(1, int(now_ts - curr_copy.get("start_ts", now_ts)))
        curr_copy["duration_sec"] = dur
        curr_copy["duration_str"] = format_duration(dur)
        curr_copy["is_current"] = True
        timeline_list.append(curr_copy)

    for item in data.get("history", []):
        timeline_list.append(item)

    is_online = (now_ts - data.get("last_seen_ts", 0)) < 15

    return JSONResponse(content={
        "base_id": base_id,
        "boot_time": data.get("boot_time", "-"),
        "is_online": is_online,
        "total_records": len(timeline_list),
        "timeline": timeline_list
    })

@app.post("/upload_screen/{client_id}")
async def upload_screen(client_id: str, title: str = "", idle_sec: int = 0, file: UploadFile = File(...)):
    file_path = os.path.join(IMAGE_DIR, f"screen_{client_id}.jpg")
    temp_path = os.path.join(IMAGE_DIR, f"temp_{client_id}.jpg")
    try:
        content = await file.read()
        with open(temp_path, "wb") as f:
            f.write(content)
        os.replace(temp_path, file_path)
        # 활성화된 창 이름 업데이트
        active_windows[client_id] = title
        
        # 타임라인 업데이트 (PC 단위)
        base_id = client_id.rsplit("-", 1)[0]
        update_client_timeline(base_id, title, idle_sec)

        return {"status": "success"}
    except: return {"status": "error"}

async def _wake_sequence(ws_list):
    try:
        # Windows 10/11 DWM 및 그래픽 드라이버 절전 해제를 위한 가상 입력 시퀀스
        for _ in range(3):
            await asyncio.sleep(0.05)
            for ws in ws_list:
                try:
                    await ws.send_json({"command": "key_press", "vk": 16}) # VK_SHIFT
                except: pass
        for ws in ws_list:
            try:
                await ws.send_json({"command": "mouse_move", "monitor": 1, "rx": 0.5, "ry": 0.5})
                await ws.send_json({"command": "mouse_move", "monitor": 1, "rx": 0.52, "ry": 0.52})
            except: pass
    except: pass

@app.post("/send_command/{target_id}")
async def send_command(target_id: str, request: Request):
    data = await request.json()
    command = data.get("command")

    # target_id가 'all'이면 모든 클라이언트에게 전송
    if target_id == "all":
        ws_list = list(active_websockets.values())
        for ws in ws_list:
            try:
                await ws.send_json(data)
            except: pass
        if command in ("monitor_on", "unlock"):
            asyncio.create_task(_wake_sequence(ws_list))
        return {"status": "success", "sent_to": list(active_websockets.keys())}

    # 특정 클라이언트에게 전송 (PC ID 기준)
    if target_id in active_websockets:
        ws = active_websockets[target_id]
        await ws.send_json(data)
        if command in ("monitor_on", "unlock"):
            asyncio.create_task(_wake_sequence([ws]))
        return {"status": "success"}

    return {"status": "error", "message": "Client not connected"}

@app.get("/get_blocklist")
async def get_blocklist():
    try:
        with open(BLOCKLIST_FILE, "r", encoding="utf-8") as f:
            return JSONResponse(content=json.load(f))
    except Exception:
        return JSONResponse(content=[])

@app.post("/set_blocklist")
async def set_blocklist(request: Request):
    data = await request.json()
    names = data.get("list", [])
    cleaned = sorted(set(str(n).strip() for n in names if str(n).strip()))
    with open(BLOCKLIST_FILE, "w", encoding="utf-8") as f:
        json.dump(cleaned, f, ensure_ascii=False)
    return {"status": "success", "list": cleaned}

@app.post("/report_processes/{client_id}")
async def report_processes(client_id: str, request: Request):
    data = await request.json()
    client_processes[client_id] = data.get("processes", [])
    return {"status": "success"}

@app.get("/get_processes/{client_id}")
async def get_processes(client_id: str):
    return JSONResponse(content=client_processes.get(client_id, []))

@app.post("/save_dashboard_capture")
async def save_dashboard_capture(request: Request):
    try:
        data = await request.json()
        filename = data.get("filename", f"대시보드_캡처_{int(time.time())}.png")
        image_data = data.get("image", "")
        if "," in image_data:
            image_data = image_data.split(",", 1)[1]
        
        file_bytes = base64.b64decode(image_data)
        file_path = os.path.join(CAPTURE_DIR, filename)
        with open(file_path, "wb") as f:
            f.write(file_bytes)
        
        print(f"[{time.strftime('%H:%M:%S')}] [CAPTURE] Saved: {filename} -> {CAPTURE_DIR}")
        return {"status": "success", "path": file_path, "dir": CAPTURE_DIR, "filename": filename}
    except Exception as e:
        print(f"[Error] save_dashboard_capture: {e}")
        return {"status": "error", "message": str(e)}

WALLPAPER_DIR = os.path.join("static", "wallpaper")
if not os.path.exists(WALLPAPER_DIR):
    os.makedirs(WALLPAPER_DIR, exist_ok=True)

@app.post("/api/wallpaper/upload")
async def upload_wallpaper(file: UploadFile = File(...)):
    try:
        os.makedirs(WALLPAPER_DIR, exist_ok=True)
        ext = os.path.splitext(file.filename)[1].lower()
        if not ext or ext not in (".jpg", ".jpeg", ".png", ".bmp"):
            ext = ".jpg"
        filename = f"current_wallpaper{ext}"
        file_path = os.path.join(WALLPAPER_DIR, filename)

        content = await file.read()
        with open(file_path, "wb") as f:
            f.write(content)

        rel_url = f"/static/wallpaper/{filename}?t={int(time.time())}"
        print(f"[{time.strftime('%H:%M:%S')}] [WALLPAPER] Uploaded: {filename} ({len(content)} bytes)")
        return {"status": "success", "url": rel_url}
    except Exception as e:
        print(f"[Error] upload_wallpaper: {e}")
        return JSONResponse(status_code=500, content={"status": "error", "message": str(e)})

@app.post("/api/wallpaper/gdrive")
async def gdrive_wallpaper(request: Request):
    try:
        data = await request.json()
        gdrive_url = data.get("url", "").strip()
        if not gdrive_url:
            return JSONResponse(status_code=400, content={"status": "error", "message": "구글 드라이브 링크를 입력해주세요."})

        m = re.search(r'(?:/file/d/|id=)([a-zA-Z0-9_-]{20,})', gdrive_url)
        file_id = m.group(1) if m else gdrive_url

        session = requests.Session()
        headers = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64)'}

        candidate_urls = [
            f"https://drive.usercontent.google.com/download?id={file_id}&export=download",
            f"https://drive.google.com/uc?export=download&id={file_id}",
            f"https://lh3.googleusercontent.com/d/{file_id}"
        ]

        content = None
        for u in candidate_urls:
            try:
                resp = session.get(u, headers=headers, timeout=15)
                if resp.status_code == 200 and len(resp.content) > 1000 and b'<!DOCTYPE html>' not in resp.content[:150]:
                    content = resp.content
                    break
            except Exception:
                continue

        if not content:
            # Fallback with confirmation cookie (for larger files)
            try:
                resp = session.get(f"https://docs.google.com/uc?export=download&id={file_id}", headers=headers, timeout=15)
                token = None
                for key, value in resp.cookies.items():
                    if key.startswith('download_warning'):
                        token = value
                        break
                if token:
                    resp = session.get(f"https://docs.google.com/uc?export=download&id={file_id}&confirm={token}", headers=headers, timeout=15)
                    if resp.status_code == 200:
                        content = resp.content
            except Exception:
                pass

        if not content or len(content) < 1000 or b'<!DOCTYPE html>' in content[:150]:
            return JSONResponse(status_code=400, content={
                "status": "error", 
                "message": "구글 드라이브에서 이미지를 가져올 수 없습니다. 링크 권한이 '링크가 있는 모든 사용자에게 공개(뷰어)'인지 확인해주세요."
            })

        os.makedirs(WALLPAPER_DIR, exist_ok=True)
        filename = "current_wallpaper.jpg"
        file_path = os.path.join(WALLPAPER_DIR, filename)
        with open(file_path, "wb") as f:
            f.write(content)

        rel_url = f"/static/wallpaper/{filename}?t={int(time.time())}"
        print(f"[{time.strftime('%H:%M:%S')}] [WALLPAPER] Downloaded from GDrive: {file_id} ({len(content)} bytes)")
        return {"status": "success", "url": rel_url}
    except Exception as e:
        print(f"[Error] gdrive_wallpaper: {e}")
        return JSONResponse(status_code=500, content={"status": "error", "message": str(e)})

# ===== 과제물 회수 및 제출 관리 =====
DEFAULT_ASSIGNMENT_BASE_DIR = os.path.join(os.getcwd(), "과제물_회수")
os.makedirs(DEFAULT_ASSIGNMENT_BASE_DIR, exist_ok=True)

assignment_state = {
    "title": "실습 과제",
    "save_dir": os.path.join(DEFAULT_ASSIGNMENT_BASE_DIR, time.strftime("%Y-%m-%d_과제")),
    "is_active": False,
    "submissions": {} # client_id -> { "files": [name, ...], "time": "16:15", "count": 2 }
}

@app.get("/api/assignment/config")
async def get_assignment_config():
    default_dir = os.path.join(DEFAULT_ASSIGNMENT_BASE_DIR, time.strftime("%Y-%m-%d_과제"))
    if not assignment_state["save_dir"]:
        assignment_state["save_dir"] = default_dir
    return JSONResponse(content={
        "title": assignment_state["title"],
        "save_dir": assignment_state["save_dir"],
        "default_dir": default_dir,
        "is_active": assignment_state["is_active"],
        "submissions": assignment_state["submissions"]
    })

@app.post("/api/assignment/start")
async def start_assignment(request: Request):
    try:
        data = await request.json()
        title = data.get("title", "").strip() or time.strftime("%Y-%m-%d 실습과제")
        save_dir = data.get("save_dir", "").strip()
        target_id = data.get("target_id", "all")

        if not save_dir:
            save_dir = os.path.join(DEFAULT_ASSIGNMENT_BASE_DIR, time.strftime("%Y-%m-%d_과제"))

        os.makedirs(save_dir, exist_ok=True)

        assignment_state["title"] = title
        assignment_state["save_dir"] = os.path.abspath(save_dir)
        assignment_state["is_active"] = True

        # WebSocket으로 학생들에게 제출창 열기 명령 전송
        cmd_payload = {"command": "open_submit_window", "title": title}
        sent_to = []
        if target_id == "all":
            for cid, ws in list(active_websockets.items()):
                try:
                    await ws.send_json(cmd_payload)
                    sent_to.append(cid)
                except Exception:
                    pass
        elif target_id in active_websockets:
            try:
                await active_websockets[target_id].send_json(cmd_payload)
                sent_to.append(target_id)
            except Exception:
                pass

        print(f"[{time.strftime('%H:%M:%S')}] [ASSIGNMENT] Started: '{title}' -> {save_dir} (Sent to {len(sent_to)} PCs)")
        return {"status": "success", "title": title, "save_dir": save_dir, "sent_to": sent_to}
    except Exception as e:
        print(f"[Error] start_assignment: {e}")
        return JSONResponse(status_code=500, content={"status": "error", "message": str(e)})

@app.post("/api/assignment/open_folder")
async def open_assignment_folder(request: Request):
    try:
        data = await request.json()
        target_path = data.get("save_dir", "").strip() or assignment_state.get("save_dir", "")
        if not target_path or not os.path.exists(target_path):
            os.makedirs(target_path or DEFAULT_ASSIGNMENT_BASE_DIR, exist_ok=True)
            target_path = target_path or DEFAULT_ASSIGNMENT_BASE_DIR

        import subprocess
        subprocess.Popen(f'explorer "{os.path.abspath(target_path)}"')
        return {"status": "success", "path": target_path}
    except Exception as e:
        return JSONResponse(status_code=500, content={"status": "error", "message": str(e)})

@app.get("/api/assignment/student_info/{client_id}")
async def get_student_info(client_id: str):
    names = load_student_names()
    name = names.get(client_id, "")
    return JSONResponse(content={
        "client_id": client_id,
        "name": name,
        "assignment_title": assignment_state.get("title", "과제물 제출")
    })

@app.post("/api/assignment/submit")
async def submit_assignment(client_id: str = Form(...), files: List[UploadFile] = File(...)):
    try:
        save_dir = assignment_state.get("save_dir")
        if not save_dir:
            save_dir = os.path.join(DEFAULT_ASSIGNMENT_BASE_DIR, time.strftime("%Y-%m-%d_과제"))
            assignment_state["save_dir"] = save_dir
        os.makedirs(save_dir, exist_ok=True)

        names = load_student_names()
        student_name = names.get(client_id, "")
        folder_name = f"{client_id}_{student_name}" if student_name else client_id
        student_folder = os.path.join(save_dir, folder_name)
        os.makedirs(student_folder, exist_ok=True)

        saved_files = []
        for file in files:
            safe_name = os.path.basename(file.filename)
            dest_file_path = os.path.join(student_folder, safe_name)
            content = await file.read()
            with open(dest_file_path, "wb") as f:
                f.write(content)
            saved_files.append(safe_name)

        submit_time = time.strftime("%H:%M:%S")
        existing = assignment_state["submissions"].get(client_id, {}).get("files", [])
        combined_files = list(set(existing + saved_files))
        assignment_state["submissions"][client_id] = {
            "files": combined_files,
            "count": len(combined_files),
            "last_time": submit_time,
            "folder": student_folder
        }

        print(f"[{submit_time}] [ASSIGNMENT SUBMISSION] {client_id} ({student_name}): {len(saved_files)} files -> {student_folder}")
        return {"status": "success", "submitted_files": saved_files, "folder": student_folder}
    except Exception as e:
        print(f"[Error] submit_assignment: {e}")
        return JSONResponse(status_code=500, content={"status": "error", "message": str(e)})

@app.get("/api/assignment/status")
async def get_assignment_status():
    return JSONResponse(content=assignment_state)

start_discovery_server()

def open_browser_when_ready():
    for _ in range(60):
        time.sleep(0.1)
        try:
            s = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
            s.settimeout(0.3)
            s.connect(("127.0.0.1", 8888))
            s.close()
            import webbrowser
            webbrowser.open("http://localhost:8888")
            break
        except Exception:
            pass

if __name__ == "__main__":
    threading.Thread(target=open_browser_when_ready, daemon=True).start()
    while True:
        try:
            if sys.platform == "win32":
                asyncio.set_event_loop_policy(asyncio.WindowsProactorEventLoopPolicy())
            uvicorn.run(app, host="0.0.0.0", port=8888, access_log=False, loop="asyncio")
        except Exception as e:
            print(f"[Server Crash Recovery] {e}. Restarting server in 1s...")
            time.sleep(1)
